"""Build the poetry archive data from the master spreadsheet.

Reads data/source/Manifesto_for_Love_MASTER.xlsx and writes:
  data/archive.json  - for tools and anything served over HTTP
  data/archive.js    - the same data as a script, so the archive pages also
                       work when opened directly from disk (file://)

Usage:
  python scripts/build_archive_data.py [path/to/spreadsheet.xlsx]
"""

import json
import re
import sys
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = ROOT / "data" / "source" / "Manifesto_for_Love_MASTER.xlsx"
# Notes shown to readers beside works that need historical context (Reference ID -> note),
# e.g. works written by slavers or from plantation slave societies
CONTEXT_NOTES = json.loads((ROOT / "data" / "context_notes.json").read_text(encoding="utf-8"))
HEADER_ROW = 6

# Spreadsheet header -> field name in the output
COLUMNS = {
    "Reference ID": "refId",
    "Country": "country",
    "Continent": "continent",
    "Work title / identifier": "title",
    "Title in English": "titleEnglish",
    "Creator / performer": "creator",
    "Open original": "openOriginal",
    "Reference type": "referenceType",
    "Form": "form",
    "Gender group": "genderGroup",
    "Year for sorting": "year",
    "Language / community": "language",
    "Tradition / repertoire": "tradition",
    "Form details": "formDetails",
    "Gender / role details": "genderDetails",
    "Original work date": "originalDate",
    "Edition / translation date": "editionDate",
    "Description / context": "description",
    "Provenance / edition": "provenance",
    "Location within source": "location",
    "Translation / version notes": "translation",
    "Original rights / reuse basis": "rights",
    "Source details": "source",
}
PASSAGE_COLUMNS = {
    "Selected passage — original language": "original",
    "Selected passage — English": "english",
}
LINK_COLUMNS = [
    "Source link",
    "Rights / catalogue link",
    "Date evidence link",
    "Creator / country evidence link",
    "Additional link 1",
    "Additional link 2",
]


def form_group(form):
    """Collapse the many specific forms into a few browsable groups."""
    f = form.lower()
    is_song = "song" in f
    is_poem = "poetry" in f or "poem" in f
    if is_song and is_poem:
        return "Poem & song"
    if is_song:
        return "Song"
    if is_poem:
        return "Poem"
    return "Narrative & oral"


def translation_status(entry, passage):
    """How the English passage relates to the original."""
    if passage["original"] == passage["english"]:
        return "english-original"
    if re.search(r"chatgpt|claude", entry["translation"], re.I):
        return "draft"
    return "translation"


# Research notes that matter to editors but not to readers
EDITORIAL = re.compile(r"confirm|verify|establish|disputed|unspecified|uncertain|unknown|not stated|attribution", re.I)


def reader_title(title):
    """Drop bracketed cataloguing notes: "(incipit)", "(first line; ...)"."""
    return re.sub(r"\s*[(\[][^)\]]*[)\]]", "", title).strip(" ,;")


def reader_text(text):
    """Keep only the parts of a ;-separated field that aren't editorial notes."""
    parts = [p.strip() for p in text.split(";")]
    return "; ".join(p for p in parts if p and not EDITORIAL.search(p))


# The site is written in UK English; the spreadsheet's own prose sometimes isn't
UK_WORDS = {
    "toward": "towards", "fulfill": "fulfil", "fulfills": "fulfils", "fulfillment": "fulfilment",
    "color": "colour", "colors": "colours", "colored": "coloured", "colorful": "colourful",
    "favor": "favour", "favorite": "favourite", "honor": "honour", "honored": "honoured",
    "behavior": "behaviour", "center": "centre", "centered": "centred", "theater": "theatre",
    "gray": "grey", "traveled": "travelled", "traveling": "travelling", "traveler": "traveller",
    "catalog": "catalogue", "jewelry": "jewellery", "defense": "defence", "skeptical": "sceptical",
    "neighbor": "neighbour", "neighbors": "neighbours", "labor": "labour", "humor": "humour",
    "rumor": "rumour", "splendor": "splendour", "ardor": "ardour", "vigor": "vigour",
    "savior": "saviour", "endeavor": "endeavour", "candor": "candour", "valor": "valour",
    "harbor": "harbour", "flavor": "flavour", "modeled": "modelled", "marveled": "marvelled",
}
IZE_WORD = re.compile(r"\b\w+(?:iz|yz)(?:e|es|ed|ing|ation|ations)\b", re.I)
NOT_IZE = {"size", "sizes", "sized", "sizing", "prize", "prizes", "prized", "seize", "seized",
           "seizes", "seizing", "capsize", "capsized", "maize", "baize", "belize", "oversized"}


def uk_spelling(text):
    def word(m):
        w = m.group(0)
        uk = UK_WORDS.get(w.lower())
        return (uk[0].upper() + uk[1:] if w[0].isupper() else uk) if uk else w

    def ize(m):
        w = m.group(0)
        if w.lower() in NOT_IZE:
            return w
        return re.sub(r"(i|y)z(?=e|ed|es|ing|ation)", r"\1s", w, count=1, flags=re.I)

    text = re.sub(r"\b[A-Za-z]+\b", word, text)
    return IZE_WORD.sub(ize, text)


def add_display_fields(entry):
    title = reader_title(entry["title"]) or entry["title"]
    english = uk_spelling(reader_title(entry["titleEnglish"]))
    entry["description"] = uk_spelling(entry["description"])
    entry["display"] = {
        "title": title,
        "english": english if english and english.lower() != title.lower() else "",
        "creator": reader_title(reader_text(entry["creator"])),
        "language": uk_spelling(reader_text(entry["language"])),
    }


def slugify(text):
    return re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")


MONTHS = "January|February|March|April|May|June|July|August|September|October|November|December"
# Tokens that legitimately run letters and digits together
KEEP_AS_IS = re.compile(r"^(10\.\d+/\S+|nlp\d+\W*|Sean2Miles\W*)$")


def restore_spaces(token):
    """The source text lost the spaces around numbers ("track6,1998",
    "of33poems", "releasedSeptember2015"). Put them back without splitting
    archive codes such as UNES08266 or C3P11, which use capitals."""
    if KEEP_AS_IS.match(token):
        return token
    token = re.sub(rf"([a-z])({MONTHS})", r"\1 \2", token)
    token = re.sub(r"([a-z'])(\d)", r"\1 \2", token)
    token = re.sub(r"(\d)([A-Z][a-z])", r"\1 \2", token)
    token = re.sub(r"(\d)(?!st|nd|rd|th)([a-z]{2,})", r"\1 \2", token)
    token = re.sub(r",(\d)", r", \1", token)
    return token


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = re.sub(r"\s+", " ", str(value)).strip()
    return " ".join(restore_spaces(t) for t in text.split(" "))


def clean_passage(value):
    """Poem text: keep the line breaks, touch nothing else."""
    if value is None:
        return ""
    lines = str(value).replace("\r\n", "\n").split("\n")
    return "\n".join(line.rstrip() for line in lines).strip()


def build(source):
    ws = openpyxl.load_workbook(source, data_only=True).worksheets[0]
    headers = [clean(c.value) for c in ws[HEADER_ROW]]
    col = {h: i for i, h in enumerate(headers)}
    missing = [h for h in [*COLUMNS, *PASSAGE_COLUMNS, *LINK_COLUMNS] if h not in col]
    if missing:
        sys.exit(f"Spreadsheet is missing expected columns: {missing}")

    entries = []
    for row in ws.iter_rows(min_row=HEADER_ROW + 1):
        if not clean(row[col["Country"]].value):
            continue
        entry = {field: clean(row[col[h]].value) for h, field in COLUMNS.items()}
        entry["id"] = slugify(entry["country"])
        if entry["refId"] in CONTEXT_NOTES:
            entry["contextNote"] = CONTEXT_NOTES[entry["refId"]]
        entry["title"] = entry["title"] or entry["titleEnglish"] or entry["tradition"]
        entry["formGroup"] = form_group(entry["form"])
        add_display_fields(entry)

        passage = {field: clean_passage(row[col[h]].value) for h, field in PASSAGE_COLUMNS.items()}
        if passage["original"] and passage["english"]:
            passage["status"] = translation_status(entry, passage)
            entry["passage"] = passage
        elif passage["original"] or passage["english"]:
            print(f"Warning: {entry['country']} has only one of the two passages; skipped")

        entry["links"] = []
        for h in LINK_COLUMNS:
            cell = row[col[h]]
            label = clean(cell.value)
            url = cell.hyperlink.target if cell.hyperlink else ""
            if url:
                if not label or label.startswith("http"):
                    label = urlparse(url).hostname.removeprefix("www.")
                entry["links"].append({"label": label, "url": url})
        entries.append(entry)

    entries.sort(key=lambda e: e["country"])
    return entries


def main():
    source = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_SOURCE
    entries = build(source)
    out = ROOT / "data"
    payload = json.dumps(entries, ensure_ascii=False, indent=2)
    (out / "archive.json").write_text(payload + "\n", encoding="utf-8")
    (out / "archive.js").write_text(
        "// Generated by scripts/build_archive_data.py - do not edit by hand.\n"
        f"window.ARCHIVE = {payload};\n",
        encoding="utf-8",
    )
    with_text = [e for e in entries if "passage" in e]
    drafts = sum(1 for e in with_text if e["passage"]["status"] == "draft")
    print(f"Wrote {len(entries)} entries to data/archive.json and data/archive.js")
    noted = [e["country"] for e in entries if "contextNote" in e]
    print(f"Context notes on {len(noted)}: {', '.join(noted)}")
    print(f"{len(with_text)} with original and English passages ({drafts} English drafts awaiting review)")
    no_url = [e["country"] for e in entries if not e["links"]]
    if no_url:
        print(f"Entries without any link: {', '.join(no_url)}")


if __name__ == "__main__":
    main()
